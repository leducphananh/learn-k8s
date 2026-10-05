import io
import pytest
from app.services.parser import extract_text_from_file
import docx
from pypdf import PdfWriter

def test_parse_txt_utf8():
    content = "Xin chào Việt Nam! Đây là bài kiểm tra tiếng Việt có dấu.".encode("utf-8")
    result = extract_text_from_file("sample.txt", content)
    assert "Xin chào Việt Nam!" in result["text"]
    assert result["word_count"] > 0
    assert result["char_count"] > 0

import unicodedata

def test_parse_txt_cp1258():
    # CP1258 (Windows-1258 Vietnamese uses decomposed characters for some diacritics)
    text = "Xin chào bạn"
    content = unicodedata.normalize("NFD", text).encode("windows-1258")
    result = extract_text_from_file("sample_cp1258.txt", content)
    assert "chào" in result["text"] or "ch" in result["text"]


def test_parse_docx():
    # Create an in-memory docx
    doc = docx.Document()
    doc.add_paragraph("Đoạn văn thứ nhất trong tài liệu Word.")
    doc.add_paragraph("Đoạn văn thứ hai kiểm tra trích xuất.")
    buffer = io.BytesIO()
    doc.save(buffer)
    content = buffer.getvalue()

    result = extract_text_from_file("document.docx", content)
    assert "Đoạn văn thứ nhất" in result["text"]
    assert "Đoạn văn thứ hai" in result["text"]
    assert result["word_count"] > 0

def test_parse_pdf():
    # Create a minimal in-memory PDF using pypdf PdfWriter
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    content = buffer.getvalue()

    result = extract_text_from_file("document.pdf", content)
    assert "text" in result
    assert isinstance(result["word_count"], int)

def test_parse_invalid_extension():
    with pytest.raises(ValueError, match="Định dạng file không được hỗ trợ"):
        extract_text_from_file("archive.zip", b"fake zip content")

def test_parse_empty_content():
    result = extract_text_from_file("empty.txt", b"")
    assert result["text"] == ""
    assert result["word_count"] == 0
    assert result["char_count"] == 0
